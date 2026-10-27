"""Excel and CSV report export studio with professional styling."""
import io
from typing import Any, Dict, List
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def export_to_csv(records: List[Dict[str, Any]]) -> bytes:
    """Generate UTF-8-SIG encoded CSV for Excel compatibility."""
    df = pd.DataFrame(records)
    # Re-order and rename columns to match the official Intelligence Portal report format
    rename_map = {
        "staff_name": "STAFF NAME",
        "department": "DEPARTMENT",
        "citations_all": "CITATIONS (ALL)",
        "citations_recent": "CITATIONS (2021+)",
        "h_index_all": "H-INDEX (ALL)",
        "h_index_recent": "H-INDEX (2021+)",
        "i10_index_all": "I10-INDEX (ALL)",
        "i10_index_recent": "I10-INDEX (2021+)",
        "match_status": "STATUS",
        "profile_name": "SCHOLAR PROFILE NAME",
        "affiliation": "AFFILIATION",
        "profile_url": "PROFILE URL"
    }
    cols = [c for c in rename_map.keys() if c in df.columns]
    out_df = df[cols].rename(columns=rename_map)
    return out_df.to_csv(index=False).encode("utf-8-sig")


def export_to_excel(
    records: List[Dict[str, Any]],
    kpis: Dict[str, Any],
    departments: List[Dict[str, Any]]
) -> bytes:
    """Generate multi-sheet styled Excel workbook."""
    wb = openpyxl.Workbook()

    # Styling presets
    navy_fill = PatternFill(start_color="1B2A4A", end_color="1B2A4A", fill_type="solid")
    header_fill = PatternFill(start_color="2A4365", end_color="2A4365", fill_type="solid")
    accent_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    navy_title = Font(name="Calibri", size=16, bold=True, color="1B2A4A")
    navy_subtitle = Font(name="Calibri", size=12, bold=True, color="2A4365")
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0")
    )

    # 1. Sheet 1: Executive Summary
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary["A1"] = "University Research Analytics — Executive Summary"
    ws_summary["A1"].font = navy_title

    # KPI Table
    kpi_rows = [
        ("Total Lecturers Ingested", kpis.get("total_staff", 0)),
        ("Distinct Departments Covered", kpis.get("unique_departments", 1)),
        ("Google Scholar Profiles Matched", kpis.get("matched_count", 0)),
        ("Match Success Rate", f"{kpis.get('match_rate', 0)}%"),
        ("Verified Profiles", kpis.get("verified_count", 0)),
        ("Needs Verification / Review", kpis.get("needs_verification_count", 0)),
        ("Total Citations (All-Time)", kpis.get("total_citations_all", 0)),
        ("Recent Citations (Recent)", kpis.get("total_citations_recent", 0)),
        ("Average Lecturer h-index", kpis.get("avg_h_index", 0.0)),
        ("Maximum Lecturer h-index", kpis.get("max_h_index", 0)),
        ("Average i10-index", kpis.get("avg_i10_index", 0.0))
    ]

    ws_summary["A3"] = "Metric"
    ws_summary["B3"] = "Value"
    ws_summary["A3"].fill = header_fill
    ws_summary["B3"].fill = header_fill
    ws_summary["A3"].font = white_bold
    ws_summary["B3"].font = white_bold

    for i, (m, v) in enumerate(kpi_rows, start=4):
        ws_summary[f"A{i}"] = m
        ws_summary[f"B{i}"] = v
        ws_summary[f"A{i}"].font = cell_font
        ws_summary[f"B{i}"].font = Font(name="Calibri", size=10, bold=True)
        ws_summary[f"A{i}"].border = thin_border
        ws_summary[f"B{i}"].border = thin_border
        if i % 2 == 0:
            ws_summary[f"A{i}"].fill = accent_fill
            ws_summary[f"B{i}"].fill = accent_fill

    # Department Leaderboard Section
    start_dept_row = len(kpi_rows) + 6
    ws_summary[f"A{start_dept_row - 1}"] = "Departmental Breakdown"
    ws_summary[f"A{start_dept_row - 1}"].font = navy_subtitle

    dept_headers = ["Department", "Staff Count", "Matched", "Total Citations", "Recent Citations", "Avg h-index", "Top Researcher"]
    for c_idx, h in enumerate(dept_headers, start=1):
        cell = ws_summary.cell(row=start_dept_row, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = white_bold

    for r_idx, d in enumerate(departments, start=start_dept_row + 1):
        vals = [
            d.get("department"), d.get("staff_count"), d.get("matched_count"),
            d.get("total_citations"), d.get("recent_citations"), d.get("avg_h_index"),
            d.get("top_researcher")
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=v)
            cell.font = cell_font
            cell.border = thin_border
            if r_idx % 2 == 0:
                cell.fill = accent_fill

    # 2. Sheet 2: Faculty Roster
    ws_roster = wb.create_sheet(title="Faculty Citation Roster")
    ws_roster.views.sheetView[0].showGridLines = True

    roster_headers = [
        "Staff Name", "Department", "Status", "Scholar Profile Name",
        "Affiliation", "Citations (All)", "Citations (Recent)",
        "h-index (All)", "h-index (Recent)", "i10-index (All)", "i10-index (Recent)", "Scholar URL"
    ]
    for c_idx, h in enumerate(roster_headers, start=1):
        cell = ws_roster.cell(row=1, column=c_idx, value=h)
        cell.fill = navy_fill
        cell.font = white_bold

    for r_idx, r in enumerate(records, start=2):
        vals = [
            r.get("staff_name", ""), r.get("department", ""), r.get("match_status", ""),
            r.get("profile_name", ""), r.get("affiliation", ""),
            int(r.get("citations_all") or 0), int(r.get("citations_recent") or 0),
            int(r.get("h_index_all") or 0), int(r.get("h_index_recent") or 0),
            int(r.get("i10_index_all") or 0), int(r.get("i10_index_recent") or 0),
            r.get("profile_url", "")
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws_roster.cell(row=r_idx, column=c_idx, value=v)
            cell.font = cell_font
            cell.border = thin_border
            if r_idx % 2 == 0:
                cell.fill = accent_fill

    # 3. Sheet 3: Verification Follow-up
    ws_verify = wb.create_sheet(title="Verification Follow-up")
    ws_verify.views.sheetView[0].showGridLines = True

    for c_idx, h in enumerate(["Staff Name", "Department", "Current Status", "Candidate Count", "Suggested Profile", "Affiliation"], start=1):
        cell = ws_verify.cell(row=1, column=c_idx, value=h)
        cell.fill = PatternFill(start_color="C53030", end_color="C53030", fill_type="solid")
        cell.font = white_bold

    unresolved = [r for r in records if str(r.get("match_status", "")).lower() != "verified"]
    for r_idx, r in enumerate(unresolved, start=2):
        vals = [
            r.get("staff_name", ""), r.get("department", ""), r.get("match_status", ""),
            r.get("candidate_count", 0), r.get("profile_name", ""), r.get("affiliation", "")
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws_verify.cell(row=r_idx, column=c_idx, value=v)
            cell.font = cell_font
            cell.border = thin_border
            if r_idx % 2 == 0:
                cell.fill = accent_fill

    # Auto-adjust column widths for all sheets
    for sheet in [ws_summary, ws_roster, ws_verify]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
